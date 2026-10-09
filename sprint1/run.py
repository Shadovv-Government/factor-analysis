"""Reconcile immutable source workbooks and produce the sprint deliverables."""
import argparse
import csv
import hashlib
import json
import sys
from copy import deepcopy
from datetime import date, datetime
from decimal import Decimal, localcontext
from pathlib import Path

from openpyxl import load_workbook
from .model import Deal, build_schedule, discount_schedule, summarize

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'outputs/sprint1_20261009'
SOURCE_FILES=('annuity-payment-schedule.xlsx','payment_schedules_discounted.xlsx')
MONEY_FIELDS={'opening_balance','interest','principal','payment','closing_balance',
    'total_interest','total_principal','total_payments','first_payment','last_payment'}


def serial(value, key=''):
    if isinstance(value,Decimal):
        if key in MONEY_FIELDS: return f'{value:.2f}'
        if 'rate' in key or key=='discount_factor': return f'{value:.18f}'
        return f'{value:.8f}'
    if isinstance(value,dict): return {k:serial(v,k) for k,v in value.items()}
    if isinstance(value,(list,tuple)): return [serial(x,key) for x in value]
    return value


def reconcile_sources(root: Path, deal: Deal, schedules: dict) -> dict:
    """Compare each cell; cache comparison is not a native Excel recalculation."""
    differences=[]; comparisons=0

    def compare(label, actual, expected, tolerance=Decimal(0)):
        nonlocal comparisons
        comparisons+=1
        if isinstance(expected,Decimal):
            try:
                observed=Decimal(str(actual))
                equal=observed.is_finite() and abs(observed-expected)<=tolerance
            except (ValueError,TypeError,ArithmeticError):
                equal=False
        elif isinstance(actual,(date,datetime)):
            equal=actual.strftime('%Y-%m-%d')==expected
        else:
            equal=actual==expected
        if not equal: differences.append({'location':label,'source_value':str(actual),'calculated':str(expected)})

    source=load_workbook(root/SOURCE_FILES[0],data_only=True,read_only=True)
    ws=source['График платежей']
    ann_rows=[r for r in ws.iter_rows(min_row=10,max_col=6,values_only=True) if isinstance(r[1],(datetime,date))]
    compare('source ANN row count',len(ann_rows),len(schedules.get('ANN',[])))
    for idx,(reference,calculated) in enumerate(zip(ann_rows,schedules.get('ANN',[])),10):
        for col,key in ((1,'date'),(2,'payment'),(3,'interest'),(4,'principal'),(5,'closing_balance')):
            compare(f'{SOURCE_FILES[0]}:График платежей:{idx}:{key}',reference[col],calculated[key])
    for coordinate,expected in [('I10',deal.valuation_date.isoformat()),('I11',deal.principal),('I12',deal.credit_rate),('I13',deal.frequency_months),('I14',deal.maturity_date.isoformat()),('I15',deal.first_payment_date.isoformat())]:
        compare(f'{SOURCE_FILES[0]}:{coordinate}',ws[coordinate].value,expected)
    source.close()

    source=load_workbook(root/SOURCE_FILES[1],data_only=True,read_only=True)
    method=source['Методика']
    for cell,expected in [('B2',deal.valuation_date.isoformat()),('B3',deal.principal),('B4',deal.credit_rate),('B5',deal.discount_rate)]:
        compare(f'{SOURCE_FILES[1]}:Методика:{cell}',method[cell].value,expected)
    fields=('number','date','days','opening_balance','interest','principal','payment','closing_balance','credit_period_rate','discount_factor','pv_interest','pv_principal','pv_payment')
    total_rows=0
    for kind in ('ANN','SIMPLE','RGM'):
        rows=list(source[kind].iter_rows(min_row=2,max_col=13,values_only=True))
        total_rows+=len(rows)
        compare(f'{kind}:row count',len(rows),len(schedules.get(kind,[])))
        for idx,(reference,calculated) in enumerate(zip(rows,schedules.get(kind,[])),2):
            for col,key in enumerate(fields):
                tol=Decimal('0.0000001') if key.startswith('pv_') else Decimal('0.000000000001') if key in ('credit_period_rate','discount_factor') else Decimal(0)
                compare(f'{SOURCE_FILES[1]}:{kind}:row {idx}:{key}',reference[col],calculated[key],tol)
        if kind in schedules:
            summary=summarize(deal,schedules[kind])
            summary_keys=('payment_count','total_payments','total_interest','total_principal','pv_future_cashflows','pv_minus_outstanding','first_payment','last_payment','closing_balance')
            summary_row={'ANN':2,'SIMPLE':3,'RGM':4}[kind]
            for col,key in enumerate(summary_keys,2):
                compare(f'{SOURCE_FILES[1]}:Сравнение:{summary_row}:{key}',source['Сравнение'].cell(summary_row,col).value,summary[key],Decimal('0.0000001') if key.startswith('pv_') else Decimal(0))
    source.close()
    return {'passed':not differences,'source_ann_rows':len(ann_rows),'discounted_rows':total_rows,
        'comparisons':comparisons,'mismatch_count':len(differences),'mismatches':differences[:20],
        'money_tolerance':'0.00','pv_tolerance':'0.00000010','factor_tolerance':'0.000000000001',
        'source_mode':'saved_values_not_native_recalculation'}


def verify_schedule(deal: Deal, rows: list[dict]) -> dict:
    failures=[]
    previous_balance=deal.principal
    for r in rows:
        if r['opening_balance']!=previous_balance: failures.append(f"row {r['number']}: balance continuity")
        if r['payment']!=r['interest']+r['principal']: failures.append(f"row {r['number']}: payment split")
        if r['closing_balance']!=r['opening_balance']-r['principal'] or r['closing_balance']<0: failures.append(f"row {r['number']}: principal roll-forward")
        with localcontext() as ctx:
            ctx.prec=40
            if abs(r['pv_payment']-r['pv_interest']-r['pv_principal'])>Decimal('1e-24'):
                failures.append(f"row {r['number']}: PV split")
        previous_balance=r['closing_balance']
    principal_diff=sum(r['principal'] for r in rows)-deal.principal
    if principal_diff != 0: failures.append('total principal')
    if rows[-1]['closing_balance']!=0: failures.append('final balance')
    return {'passed':not failures,'checked_payments':len(rows),'principal_difference':f'{principal_diff:.2f}',
        'closing_balance':f"{rows[-1]['closing_balance']:.2f}",'failures':failures}


def source_hashes(root: Path) -> list[dict]:
    return [{'file':name,'sha256':hashlib.sha256((root/name).read_bytes()).hexdigest()} for name in SOURCE_FILES]


def _rub(value):
    return f'{Decimal(str(value)):,.2f}'.replace(',',' ').replace('.',',')


def report(payload,verification):
    if verification['references'].get('passed') is not True:
        return generic_report(payload, verification)
    req=payload['request']; deal=req['deal']; results=payload['results']
    lines=[
        '# Денежные потоки сделки и результаты первого спринта', '',
        '## Результат', '',
        'Подготовлен воспроизводимый прототип ANN, SIMPLE и RGM. Расчёт выполняется на Python с Decimal; новый Excel показывает формулы. Все суммы ниже относятся к будущим потокам после даты оценки. Это расчёт одной сделки, который станет основой последующих этапов факторного анализа.', '',
        '## Исходные данные и допущения', '',
        f"Долг {_rub(deal['outstanding_principal'])} RUB на {deal['valuation_date']}; кредитная ставка {Decimal(deal['annual_credit_rate'])*100}% годовых. Первый платёж {deal['first_payment_date']}, погашение {deal['maturity_date']}, интервал {deal['frequency_months']} месяц(ев). Дата открытия 22.06.2020 сохранена в исходном файле справочно: прототип начинается с известного остатка, прошлые выплаты не восстанавливаются.", '',
        f"Дисконтная ставка задаётся отдельно: {Decimal(req['discounting']['annual_rate'])*100}% годовых в базовом примере. Это учебное допущение; рыночная кривая не предоставлена. ACT/365F делит фактическое число дней на 365 и в високосном году. Выходные не сдвигают даты. Комиссии, страховки, досрочные погашения, дефолты и проценты, накопленные до даты оценки, не заданы.", '',
        'Проценты и погашение тела округляются HALF_UP до двух знаков. Коэффициенты дисконтирования и PV вычисляются с 40 значащими цифрами без промежуточного округления; в JSON и CSV DF/ставки отображаются с 18 знаками, PV с 8. В Excel применяется ROUND(x,2).', '',
        '## Разбор исходного ANN', '',
        'В исходной книге лист «График платежей», H9:I20, содержит параметры; B10:F609 — формулы графика (334 заполненных платежа). Лист «Расчет платежа» вычисляет коэффициенты на фактических датах. ANN использует эти коэффициенты, а не ставку 7,5% / 12 для каждого интервала в выбранном режиме «Факт / 365».', '',
        'Для периода i: d_i = date_i − date_(i−1); q_i = r_credit × d_i / 365; D_0 = 1; D_i = D_(i−1) / (1 + q_i). Регулярный аннуитет A = ROUND(P_0 / ΣD_i, 2). В исходном примере ΣD_i = 140,2149596068776 и A = 26 003,20 RUB.', '',
        'Проценты I_i = ROUND(B_(i−1) × q_i, 2); платёж C_i = MIN(A, B_(i−1) + I_i); погашение тела T_i = C_i − I_i; остаток B_i = B_(i−1) − T_i. В последнюю дату C_n = B_(n−1) + I_n: это закрывает копеечный остаток после округления.', '',
        'Первый интервал 31.08.2022 → 22.09.2022 составляет 22 дня: I_1 = ROUND(3 646 037,89 × 0,075 × 22 / 365, 2) = 16 482,09; T_1 = 26 003,20 − 16 482,09 = 9 521,11; B_1 = 3 636 516,78 RUB.', '',
        'Второй интервал составляет 30 дней: I_2 = ROUND(3 636 516,78 × 0,075 × 30 / 365, 2) = 22 416,88; T_2 = 3 586,32; B_2 = 3 632 930,46 RUB.', '',
        'Последний интервал 22.05.2050 → 15.06.2050 составляет 24 дня: I_n = ROUND(25 877,84 × 0,075 × 24 / 365, 2) = 127,62; C_n = 25 877,84 + 127,62 = 26 005,46; конечный остаток 0,00 RUB.', '',
        '| № | Дата | Дней | Платёж RUB | Проценты RUB | Тело RUB | Остаток RUB |',
        '|---:|---|---:|---:|---:|---:|---:|'
    ]
    ann=results.get('ANN',{}).get('payments',[])
    for r in ann:
        if r['number'] in (1,2,3,333,334) or r['date']=='2024-03-22':
            lines.append(f"| {r['number']} | {r['date']} | {r['days']} | {_rub(r['payment'])} | {_rub(r['interest'])} | {_rub(r['principal'])} | {_rub(r['closing_balance'])} |")
    lines += ['', 'Интервал 22.02.2024 → 22.03.2024 содержит 29 дней; база остаётся 365. Даты следующих платежей привязаны к первоначальному дню месяца, последний короткий период заканчивается точно в дату погашения.', '',
        '## SIMPLE и RGM', '',
        'SIMPLE: каждый период оплачиваются проценты на неизменный долг; всё тело 3 646 037,89 RUB возвращается в последний день. Последний платёж 3 664 018,35 RUB включает 17 980,46 RUB процентов за 24 дня.', '',
        'RGM: регулярное погашение тела ROUND(3 646 037,89 / 334, 2) = 10 916,28 RUB. Последний платёж закрывает 10 916,65 RUB тела и 53,84 RUB процентов. При маленьком долге погашение ограничивается текущим остатком. Общая тенденция процентной части снижается, но фактическое число дней вызывает колебания между соседними платежами.', '',
        '## Сравнение', '',
        '| Тип | Платежей | Всего выплат RUB | Проценты RUB | Тело RUB | PV RUB | PV минус долг RUB | Конечный остаток RUB |',
        '|---|---:|---:|---:|---:|---:|---:|---:|'
    ]
    for k,r in results.items():
        s=r['summary']; lines.append(f"| {k} | {s['payment_count']} | {_rub(s['total_payments'])} | {_rub(s['total_interest'])} | {_rub(s['total_principal'])} | {_rub(s['pv_future_cashflows'])} | {Decimal(s['pv_minus_outstanding']):.8f} | {_rub(s['closing_balance'])} |")
    lines += ['', '![Остаток долга](figures/balances.png)', '',
        '![Платежи до последнего](figures/regular_payments.png)', '',
        'На графике регулярных платежей последний платёж исключён для читаемости. Крупный возврат тела SIMPLE показан числом в сравнительной таблице и в полном CSV/Excel.', '',
        '![Всего процентов](figures/interest_totals.png)', '',
        '## Дисконтирование и интерпретация', '',
        'F_0 = 1; F_i = F_(i−1) / (1 + r_discount × d_i / 365). PV_i = C_i × F_i; PV_future = ΣPV_i. PV процентов и тела считается отдельно и суммируется. PERIOD_SIMPLE_CHAIN — последовательное дисконтирование по простой ставке интервала; оно отличается от (1 + r)^(−days/365) с эффективной годовой ставкой и от обычного XNPV.', '',
        'При r_discount = r_credit без округления C_i = B_(i−1)(1+q_i) − B_i. После умножения на F_i сумма телескопируется: ΣC_iF_i = B_0 − B_nF_n. При B_n=0 PV равна начальному долгу независимо от способа погашения. Копеечное округление процентов и платежей объясняет небольшие остаточные различия в данном примере.', '',
        'Следовательно, различия номинальных процентов между графиками не доказывают разницу приведённой стоимости при этой дисконтной ставке. Для рыночной оценки нужна независимо обоснованная ставка или кривая. PV минус текущий долг — диагностическое сравнение с балансовым остатком; полный NPV от выдачи требует первоначального исходящего потока и всех прочих потоков. Economic Capital не рассчитывается на этом этапе.', '',
        '## Проверки и воспроизводимость', '',
        f"Сверка источников: {verification['references'].get('comparisons',0)} сравнений; {verification['references'].get('mismatch_count','не выполнялась')} расхождений. Сопоставлены {verification['references'].get('source_ann_rows',0)} платежа исходного ANN и {verification['references'].get('discounted_rows',0)} строки ANN/SIMPLE/RGM второго файла. Сверяются сохранённые значения источников; исходные книги не пересчитывались в установленном Excel.", '',
        'Для каждой строки проверены непрерывность долга, сумма процентов и тела, движение остатка и разделение PV. Для каждого графика сумма тела равна начальному долгу и конечный остаток нулевой. Допуски сверки: деньги 0,00 RUB; PV 0,00000010 RUB; DF и ставки периода 10^−12.', '',
        'Автотесты дополнительно проверяют нулевую ставку, ACT/365F в високосном году, календарь конца месяца, короткий финальный период, половинное округление, маленький долг, недопустимые входы и независимость ставки дисконтирования. Тесты сверки специально меняют сумму и дату, чтобы убедиться в обнаружении ошибок.', '',
        'Повторение расчёта: `python -m unittest discover -s tests -v`, затем `python -m sprint1.run`. Вход — `contracts/request.example.json`. Для другой сделки используются `--input` и `--skip-reference`; сверка с эталоном явно помечается пропущенной. CSV и JSON содержат полный график; `verification.json` и `workbook_verification.json` сохраняют результаты проверок. Команды сборки Excel и проверки интерфейса приведены в README.', '',
        '## Контракт и границы модуля', '',
        'Запрос передаёт версию, параметры сделки, отдельную ставку дисконтирования и список типов платежей. Ответ возвращает полный график, составляющие PV, итоги, допущения и источники. Деньги и ставки — десятичные строки; даты — ISO 8601. JSON Schema лежат в `contracts/`, полный ответ — `calculation.json`.', '',
        'Предложение C++: validate → build_payment_dates → build_schedule → discount → summarize → check. Ядро не читает Excel, не обращается к LLM, не выполняет факторный анализ или расчёт EC. Адаптеры занимаются JSON/Excel, командная аналитика — риск-параметрами, факторным разложением и объяснениями. Подробности: `docs/cpp-module.md`; типы и сигнатуры: `cpp/cashflows.hpp`. Это проект интерфейса; исполняемый прототип написан на Python.', '',
        '## Выполнение задач спринта', '',
        '| Задача | Подготовленные материалы | Критерий |', '|---|---|---|',
        '| Разобрать ANN | Формулы и примеры выше; сверка каждой строки | Значения источника объясняются и воспроизводятся |',
        '| Рассчитать SIMPLE и RGM | CSV, JSON, Excel; расчётный код | Сумма тела равна долгу, конечный остаток нулевой |',
        '| Дисконтировать потоки | Метод, DF/PV каждой строки, сравнение | Дата, ставка, метод и округление зафиксированы |',
        '| Описать входы и выходы | Схемы запроса/ответа и примеры JSON | Поля и ограничения имеют явные определения |',
        '| Предложить границы C++ | Документ архитектуры, header, каталог тестов | Ответственность ядра и соседних команд разделена |', '',
        '## Ограничения и следующий этап', '',
        'Версия 1 поддерживает RUB, фиксированную неотрицательную номинальную ставку, ACT/365F, от 1 до 1200 платежей и частоту 1–120 месяцев. ANN отклоняется, если регулярный платёж меньше процентов: отрицательная амортизация не поддерживается. Календарь для данной сделки фиксируется входными датами; изменение календаря требует нового запуска сборки материалов. Входные денежные суммы задаются с двумя знаками и максимум 15 цифрами до точки; десятичные строки ставок ограничены 32 символами.', '',
        'Для универсального факторного анализа потребуются наблюдения за несколькими периодами, портфель и дополнительные факторы. По одной сделке нельзя подтвердить устойчивость ML-объяснений изменения NPV. В следующем этапе можно добавить рыночную кривую, портфельную агрегацию, риск-параметры и затем факторное разложение с LLM-интерпретацией.', '',
        '## Источники', ''
    ]
    for s in payload['sources']: lines.append(f"- `{s['file']}`; SHA-256 `{s['sha256']}`.")
    lines.append('- Документ «Постановка задачи Часть 3 факторный анализ Economic Capital» и изображения «что есть», «что надо сделать»: общий контекст и пять критериев текущего спринта.')
    return '\n'.join(lines)+'\n'


def generic_report(payload, verification):
    """A different request must not inherit the source deal's worked examples."""
    d=payload['request']['deal']; disc=payload['request']['discounting']
    lines=['# Расчёт денежных потоков заданной сделки','',
        f"Сделка {d['id']}: долг {_rub(d['outstanding_principal'])} RUB на {d['valuation_date']}; кредитная ставка {d['annual_credit_rate']}; дисконтная ставка {disc['annual_rate']}. Первый платёж {d['first_payment_date']}, погашение {d['maturity_date']}.",'',
        'ACT/365F; HALF_UP до копеек. Проценты начисляются на остаток за фактические дни. ANN использует сумму последовательных кредитных DF, SIMPLE возвращает тело в конце, RGM гасит равные части тела. Последний платёж закрывает остаток. Дисконтирование выполняется отдельно по PERIOD_SIMPLE_CHAIN. PV относится к будущим потокам, а разница с остатком не является полным NPV от выдачи или EC.','',
        '| Тип | Платежей | Всего выплат RUB | Проценты RUB | Тело RUB | PV RUB | Остаток RUB |','|---|---:|---:|---:|---:|---:|---:|']
    for k,v in payload['results'].items():
        s=v['summary']
        lines.append(f"| {k} | {s['payment_count']} | {_rub(s['total_payments'])} | {_rub(s['total_interest'])} | {_rub(s['total_principal'])} | {_rub(s['pv_future_cashflows'])} | {_rub(s['closing_balance'])} |")
    lines += ['', '## Примеры платежей', '']
    for k,v in payload['results'].items():
        for label,row in [('Первый',v['payments'][0]),('Последний',v['payments'][-1])]:
            lines.append(f"{k}, {label.lower()} платёж {row['date']}: {row['days']} дней; проценты {_rub(row['interest'])}; тело {_rub(row['principal'])}; платёж {_rub(row['payment'])}; остаток {_rub(row['closing_balance'])} RUB.")
    lines += ['', '## Проверки','', 'Внутренние проверки платежей пройдены. Сверка с исходными Excel явно пропущена для данного запроса; совпадение с исходной сделкой не утверждается. Полные значения и проверки находятся в calculation.json, verification.json и CSV. Формулы Excel и изображения требуют отдельного запуска сборки для этой сделки.','']
    return '\n'.join(lines)


def validation_cases(request):
    cases=[]
    for name,block,key,value,cell in [
        ('discount_10','discounting','annual_rate','0.10','C8'),
        ('discount_0','discounting','annual_rate','0','C8'),
        ('credit_10','deal','annual_credit_rate','0.10','C7'),
        ('credit_0','deal','annual_credit_rate','0','C7'),
        ('credit_20','deal','annual_credit_rate','0.20','C7'),
        ('principal_double','deal','outstanding_principal',f"{Decimal(request['deal']['outstanding_principal'])*2:.2f}",'C6')]:
        data=deepcopy(request); data[block][key]=value
        try:
            deal=Deal.from_dict(data)
        except ValueError as exc:
            cases.append({'name':name,'cell':cell,'value':float(value),
                'status':'not_applicable','reason':str(exc),'results':{}})
            continue
        results={}
        for kind in deal.payment_types:
            try:
                rows=discount_schedule(deal,build_schedule(deal,kind))
            except ValueError as exc:
                results[kind]={'unsupported':True,'reason':str(exc)}
                continue
            results[kind]=serial({'summary':summarize(deal,rows),'first':rows[0],'last':rows[-1]})
        cases.append({'name':name,'cell':cell,'value':float(value),
            'status':'evaluated','results':results})
    return cases


def run(request: dict, source_root: Path, output: Path, *, verify_references=True) -> dict:
    deal=Deal.from_dict(request)
    schedules={k:discount_schedule(deal,build_schedule(deal,k)) for k in deal.payment_types}
    checks={k:verify_schedule(deal,rows) for k,rows in schedules.items()}
    references=reconcile_sources(source_root,deal,schedules) if verify_references else {'status':'skipped','passed':None,'reason':'Explicit --skip-reference for a different request'}
    if any(not c['passed'] for c in checks.values()) or references.get('passed') is False:
        raise ValueError(f'Verification failed: {checks}; references={references}')
    sources=source_hashes(source_root) if verify_references else []
    verification={'schedules':checks,'references':references,'sources':sources,'passed':True}
    payload=serial({'schema_version':'1.0','request':request,
        'methodology':{'day_count':'ACT/365F','discount_method':'PERIOD_SIMPLE_CHAIN',
            'rounding':'HALF_UP_2DP','calculation_precision':40,'pv_definition':'Future cashflows at valuation date',
            'assumptions':['Fixed nonnegative credit and discount rates','No holiday adjustment, fees, pre-valuation accrued interest, prepayment or default','No market discount curve provided','PV minus outstanding is not full-life NPV or Economic Capital']},
        'results':{k:{'summary':summarize(deal,rows),'payments':rows} for k,rows in schedules.items()},
        'sources':sources})
    auxiliary=validation_cases(request)
    output.mkdir(parents=True,exist_ok=True)
    (output/'calculation.json').write_text(json.dumps(payload,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    (output/'verification.json').write_text(json.dumps(verification,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    (output/'validation_cases.json').write_text(json.dumps(auxiliary,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    for kind,content in payload['results'].items():
        with (output/f'{kind}.csv').open('w',encoding='utf-8-sig',newline='') as f:
            writer=csv.DictWriter(f,fieldnames=list(content['payments'][0])); writer.writeheader(); writer.writerows(content['payments'])
    with (output/'comparison.csv').open('w',encoding='utf-8-sig',newline='') as f:
        writer=csv.DictWriter(f,fieldnames=['payment_type']+list(next(iter(payload['results'].values()))['summary'])); writer.writeheader()
        writer.writerows({'payment_type':k,**v['summary']} for k,v in payload['results'].items())
    (output/'report.md').write_text(report(payload,verification),encoding='utf-8')
    return payload


def main():
    parser=argparse.ArgumentParser(description='Спринт 1: графики, PV и сверка источников')
    parser.add_argument('--input',type=Path,default=ROOT/'contracts/request.example.json')
    parser.add_argument('--source-root',type=Path,default=ROOT)
    parser.add_argument('--output',type=Path,default=OUT)
    parser.add_argument('--skip-reference',action='store_true',help='Явно пропустить эталон для другой сделки')
    args=parser.parse_args()
    try:
        payload=run(json.loads(args.input.read_text(encoding='utf-8')),args.source_root,args.output,verify_references=not args.skip_reference)
    except (ValueError,OSError,KeyError) as exc:
        print(f'ERROR: {exc}',file=sys.stderr); return 1
    print(f'Output: {args.output}')
    for kind,result in payload['results'].items():
        s=result['summary']; print(f"{kind}: n={s['payment_count']}, interest={s['total_interest']}, PV={s['pv_future_cashflows']}, balance={s['closing_balance']}")
    return 0


if __name__=='__main__':
    raise SystemExit(main())
